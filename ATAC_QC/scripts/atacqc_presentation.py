#!/usr/bin/env python3
"""
Generate a PowerPoint QC presentation from ATAC QC pipeline outputs.

Supports two input modes:

  Local (pipeline config YAML):
    python atacqc_presentation.py config.yaml --runtag v0

  GCS (Terra / cloud):
    python atacqc_presentation.py --gcs gs://bucket/atacqc/Tissue_v0/
"""

import argparse
import os
import re
import shutil
import subprocess
import tempfile
import glob as globmod

import pandas as pd

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from PIL import Image
except ImportError as e:
    raise SystemExit(
        "Error: python-pptx and its dependencies (including Pillow) are required.\n"
        "    pip install python-pptx Pillow\n"
        f"Original error: {e}"
    )

# Optional: PDF-to-image conversion
try:
    from pdf2image import convert_from_path
    HAS_PDF2IMAGE = True
except ImportError:
    HAS_PDF2IMAGE = False

# Import pipeline utilities when available (for local mode)
try:
    from atac_qc.utils import load_config, standardize_tissue_name
except ImportError:
    try:
        import yaml
    except ImportError:
        raise SystemExit(
            "Error: PyYAML is required for standalone mode.\n"
            "    pip install pyyaml"
        )

    def load_config(config_path):
        if not os.path.exists(config_path):
            raise FileNotFoundError(f"Config file not found: {config_path}")
        with open(config_path, "r") as f:
            return yaml.safe_load(f)

    def standardize_tissue_name(tissue):
        tissue = re.sub(r'[()]', '', tissue)
        tissue = tissue.replace('-', '_')
        tissue = re.sub(r'\s+', '_', tissue)
        tissue = re.sub(r'_+', '_', tissue)
        return tissue.strip('_')


# ---------------------------------------------------------------------------
# Slide dimensions (widescreen 16:9)
# ---------------------------------------------------------------------------
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


# ---------------------------------------------------------------------------
# Sample-ID helpers
# ---------------------------------------------------------------------------

def is_per_sample(filepath, sample_ids):
    """Check if a file is per-sample (contains a sample ID in its name)."""
    base = os.path.basename(filepath)
    return any(sid in base for sid in sample_ids)


def extract_sample_ids_from_config(config, tissue_std):
    """Extract sample IDs (atacID column) from the pipeline's sample_metadata TSV."""
    meta_path = config['paths']['sample_metadata']
    tissue_config = config['params'].get('tissue', '---')
    df = pd.read_csv(meta_path, sep='\t', header=None, index_col=False,
                     names=["rnaID", "atacID", "species", "donorID",
                            "ageGroup", "gender", "tissue"])
    if tissue_config != "---":
        df = df[
            df["tissue"].apply(lambda x: standardize_tissue_name(str(x))) == tissue_std
        ]
    return set(df['atacID'].dropna().tolist())


def extract_sample_ids_from_files(tmpdir, tissue_std):
    """Infer sample IDs from downloaded filenames (GCS fallback).

    Looks for ATAC sample IDs in filenames — typically long IDs like
    RHGTEX013-603BD-01-SM-OLJCX-EXP01-ATAC-01.
    """
    sample_ids = set()
    # Match ATAC sample IDs (contain -ATAC- in the name)
    pattern = re.compile(r'([\w]+-[\w]+-[\w]+-SM-[\w]+-[\w]+-ATAC-[\w]+)')
    for root, _dirs, files in os.walk(tmpdir):
        for f in files:
            m = pattern.search(f)
            if m:
                sample_ids.add(m.group(1))
    return sample_ids


# ---------------------------------------------------------------------------
# GCS helpers
# ---------------------------------------------------------------------------

def gsutil_ls(gcs_path):
    """List GCS path contents."""
    try:
        result = subprocess.run(
            ["gsutil", "ls", gcs_path],
            capture_output=True, text=True
        )
    except FileNotFoundError as e:
        raise RuntimeError(
            "Failed to run 'gsutil ls'. Ensure gsutil is installed and on your PATH."
        ) from e

    if result.returncode != 0:
        stderr = result.stderr.strip()
        msg = f"'gsutil ls' failed for '{gcs_path}' (exit code {result.returncode})."
        if stderr:
            msg += f" stderr: {stderr}"
        raise RuntimeError(msg)

    return [line.strip() for line in result.stdout.strip().split("\n") if line.strip()]


def download_figures_and_stats(gcs_base, tmpdir):
    """Download all figures and stats from an ATAC QC GCS directory."""
    steps = {
        "1_pcr_chimera":                ["*.txt"],
        "2_load_fragments/figures":     ["*.pdf"],
        "3_qc_filter/figures":          ["*.pdf"],
        "4_doublet_detection/figures":  ["*.png"],
        "5_downstream/figures":         ["*.png", "*.pdf"],
        "5_downstream":                 ["*.tsv"],
        "6_integration/figures":        ["*.png"],
    }
    for step, exts in steps.items():
        step_dir = os.path.join(tmpdir, step.replace("/", "_"))
        os.makedirs(step_dir, exist_ok=True)
        for ext in exts:
            result = subprocess.run(
                ["gsutil", "-m", "-q", "cp",
                 gcs_base + step + "/" + ext, step_dir + "/"],
                capture_output=True, text=True
            )
            if result.returncode != 0:
                print(f"[WARNING] gsutil cp failed for {gcs_base}{step}/{ext} (may not exist yet)")
    return tmpdir


# ---------------------------------------------------------------------------
# Local directory collection
# ---------------------------------------------------------------------------

def collect_local_figures(workdir, runtag, tmpdir):
    """Copy local pipeline outputs into a normalized tmpdir structure."""
    dir_mapping = {
        f"qc_filtering.{runtag}": "3_qc_filter_figures",
        f"doublet_detection.{runtag}": "4_doublet_detection_figures",
        f"doublet_filter_processing.{runtag}": "5_downstream_figures",
        f"postprocessing_check.{runtag}": "6_integration_figures",
    }
    # Fragment plots are in fragment/ (no runtag)
    fragment_src = os.path.join(workdir, "fragment")
    if os.path.isdir(fragment_src):
        dst = os.path.join(tmpdir, "2_load_fragments_figures")
        shutil.copytree(fragment_src, dst, dirs_exist_ok=True)
        print(f"[INFO]   fragment/ -> 2_load_fragments_figures/")

    found_any = False
    for local_name, tmp_name in dir_mapping.items():
        src = os.path.join(workdir, local_name)
        if os.path.isdir(src):
            dst = os.path.join(tmpdir, tmp_name)
            shutil.copytree(src, dst, dirs_exist_ok=True)
            found_any = True
            print(f"[INFO]   {local_name}/ -> {tmp_name}/")
    if not found_any:
        print(f"[WARN] No pipeline output directories found in {workdir}")
    return tmpdir


# ---------------------------------------------------------------------------
# PDF handling
# ---------------------------------------------------------------------------

def pdf_to_images(pdf_path, output_dir=None):
    """Convert a PDF to a list of PNG image paths.

    Uses pdf2image if available, otherwise returns empty list with a warning.
    """
    if not HAS_PDF2IMAGE:
        print(f"[WARNING] pdf2image not installed — skipping PDF: {os.path.basename(pdf_path)}")
        print("          Install with: pip install pdf2image (requires poppler)")
        return []

    if output_dir is None:
        output_dir = os.path.dirname(pdf_path)

    base = os.path.splitext(os.path.basename(pdf_path))[0]
    try:
        images = convert_from_path(pdf_path, dpi=200)
    except Exception as e:
        print(f"[WARNING] Failed to convert PDF {os.path.basename(pdf_path)}: {e}")
        return []

    paths = []
    for i, img in enumerate(images):
        png_path = os.path.join(output_dir, f"{base}_page{i+1}.png")
        img.save(png_path, "PNG")
        paths.append(png_path)
    return paths


# ---------------------------------------------------------------------------
# File finding / naming helpers
# ---------------------------------------------------------------------------

def find_files(base_dir, pattern):
    """Find files matching pattern recursively, deduplicating by filename."""
    matches = []
    seen = set()
    for root, dirs, files in os.walk(base_dir):
        for f in files:
            if globmod.fnmatch.fnmatch(f, pattern) and f not in seen:
                matches.append(os.path.join(root, f))
                seen.add(f)
    matches.sort()
    return matches


def short_name(filepath):
    """Extract a readable name from the file path."""
    name = os.path.basename(filepath)
    for ext in (".png", ".pdf", ".jpg"):
        name = name.replace(ext, "")
    return name


def group_by_sample(files, sample_ids):
    """Separate files into overview and per-sample lists."""
    overview = []
    per_sample = {}
    for f in files:
        base = os.path.basename(f)
        matched_id = None
        for sid in sample_ids:
            if sid in base:
                matched_id = sid
                break
        if matched_id:
            per_sample.setdefault(matched_id, []).append(f)
        else:
            overview.append(f)
    return overview, per_sample


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def add_title_slide(prs, tissue_name, source_label):
    """Add title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"ATAC QC Report: {tissue_name.replace('_', ' ')}"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    p.alignment = PP_ALIGN.CENTER
    if source_label:
        p2 = tf.add_paragraph()
        p2.text = f"Source: {source_label}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p2.alignment = PP_ALIGN.CENTER


def add_section_slide(prs, title, subtitle=""):
    """Add a section divider slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    p.alignment = PP_ALIGN.CENTER
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(16)
        p2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p2.alignment = PP_ALIGN.CENTER


def add_image_slide(prs, image_path, title="", max_height=Inches(6.2)):
    """Add a slide with a single image, scaled to fit."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if title:
        txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(12.5), Inches(0.6))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

    with Image.open(image_path) as img:
        img_w, img_h = img.size

    max_width = Inches(12.5)
    top_margin = Inches(0.8) if title else Inches(0.3)

    scale_w = max_width / Emu(int(img_w * 914400 / 96))
    scale_h = max_height / Emu(int(img_h * 914400 / 96))
    scale = min(scale_w, scale_h, 1.0)

    final_w = int(img_w * 914400 / 96 * scale)
    final_h = int(img_h * 914400 / 96 * scale)

    left = int((SLIDE_WIDTH - final_w) / 2)
    slide.shapes.add_picture(image_path, left, int(top_margin), final_w, final_h)


def add_two_images_slide(prs, img1, img2, title=""):
    """Add a slide with two images side by side."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if title:
        txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(12.5), Inches(0.6))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

    top = Inches(0.8) if title else Inches(0.3)
    max_w = Inches(6.2)
    max_h = Inches(6.0)

    for i, img_path in enumerate([img1, img2]):
        with Image.open(img_path) as img:
            img_w, img_h = img.size
        scale_w = max_w / Emu(int(img_w * 914400 / 96))
        scale_h = max_h / Emu(int(img_h * 914400 / 96))
        scale = min(scale_w, scale_h, 1.0)
        final_w = int(img_w * 914400 / 96 * scale)
        final_h = int(img_h * 914400 / 96 * scale)
        left = Inches(0.3) if i == 0 else Inches(6.8)
        slide.shapes.add_picture(img_path, int(left), int(top), final_w, final_h)


def add_text_slide(prs, title, text_content):
    """Add a slide with text content (for stats files)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(12), Inches(6))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    lines = text_content.strip().split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p2 = tf2.paragraphs[0]
        else:
            p2 = tf2.add_paragraph()
        p2.text = line
        p2.font.size = Pt(12)
        p2.font.name = "Courier New"


def add_tsv_slide(prs, title, tsv_path):
    """Add a slide with TSV data rendered as a table."""
    try:
        df = pd.read_csv(tsv_path, sep='\t')
        text = df.to_string(index=False)
    except Exception:
        with open(tsv_path) as f:
            text = f.read()
    add_text_slide(prs, title, text)


# ---------------------------------------------------------------------------
# Main presentation builder
# ---------------------------------------------------------------------------

def build_presentation(tmpdir, tissue_name, tissue_std, sample_ids, source_label=""):
    """Build the PowerPoint presentation from figures in tmpdir."""
    print(f"[INFO] Building presentation for {tissue_name}...")
    if sample_ids:
        print(f"[INFO]   {len(sample_ids)} sample IDs for per-sample detection")
    else:
        print("[WARN]   No sample IDs found — all plots will be treated as overview")

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    add_title_slide(prs, tissue_std, source_label)

    # --- Step 1: PCR Chimera Stats ---
    pcr_stats = find_files(tmpdir, "*rmPCRchimeric_stat.txt")
    if pcr_stats:
        add_section_slide(prs, "Step 1: PCR Chimera Removal", "Read filtering statistics")
        for stat_file in pcr_stats:
            sid = os.path.basename(stat_file).split(".")[0]
            with open(stat_file) as f:
                add_text_slide(prs, f"PCR Chimera Stats: {sid}", f.read())

    # --- Step 2: Fragment Size Distribution ---
    frag_pdfs = find_files(tmpdir, "*.fragment_size_distribution.pdf")
    if frag_pdfs:
        add_section_slide(prs, "Step 2: Fragment Size Distribution",
                          "Fragment size distribution per sample")
        for pdf_path in frag_pdfs:
            pages = pdf_to_images(pdf_path, tmpdir)
            if pages:
                for page in pages:
                    add_image_slide(prs, page, short_name(pdf_path))

    # --- Step 3: QC KDE Filtering ---
    kde_before = find_files(tmpdir, "ATAC_beforeFilter_kde.*.pdf")
    kde_after = find_files(tmpdir, "ATAC_postFilter_kde.*.pdf")

    if kde_before or kde_after:
        add_section_slide(prs, "Step 3: QC Filtering (KDE)",
                          "Fragment count vs TSS enrichment")

        # Before filter PDF (multi-page, one per sample)
        for pdf_path in kde_before:
            pages = pdf_to_images(pdf_path, tmpdir)
            if pages:
                add_section_slide(prs, "Pre-Filter KDE Plots",
                                  "With cutoff lines shown")
                for page in pages:
                    add_image_slide(prs, page, "Pre-Filter KDE")

        # After filter PDF (multi-page, one per sample)
        for pdf_path in kde_after:
            pages = pdf_to_images(pdf_path, tmpdir)
            if pages:
                add_section_slide(prs, "Post-Filter KDE Plots")
                for page in pages:
                    add_image_slide(prs, page, "Post-Filter KDE")

    # --- Step 4: Doublet Detection ---
    doublet_pngs = find_files(tmpdir, "ATAC_QC_doubletHist.*.png")
    if doublet_pngs:
        add_section_slide(prs, "Step 4: Doublet Detection",
                          "Scrublet doublet score and probability distributions")
        overview, per_sample = group_by_sample(doublet_pngs, sample_ids)
        for png in overview:
            add_image_slide(prs, png, short_name(png))
        for sid in sorted(per_sample):
            for png in per_sample[sid]:
                add_image_slide(prs, png, f"Doublet Detection: {sid}")

    # --- Step 5: Downstream Processing ---
    numcell_tsv = find_files(tmpdir, "ATAC_NumCell.*.tsv")
    numcell_png = find_files(tmpdir, "ATAC_NumCell.*.png")
    cluster_pdf = find_files(tmpdir, "ATAC_post_filter_clusters.*.pdf")

    if any([numcell_tsv, numcell_png, cluster_pdf]):
        add_section_slide(prs, "Step 5: Downstream Processing",
                          "Doublet filtering, embedding, and clustering")

        # Cell count summary table
        for tsv in numcell_tsv:
            add_tsv_slide(prs, "Cell Counts per Sample", tsv)

        # Cell count histogram
        for png in numcell_png:
            add_image_slide(prs, png, "Cell Count Distribution")

        # Per-sample UMAP clusters PDF
        for pdf_path in cluster_pdf:
            pages = pdf_to_images(pdf_path, tmpdir)
            if pages:
                for page in pages:
                    add_image_slide(prs, page, "Per-Sample UMAP Clusters")

    # --- Step 6: Integration ---
    umap_sample = find_files(tmpdir, "ATAC_UMAP_bySample.*.png")
    umap_donor = find_files(tmpdir, "ATAC_UMAP_byDonor.*.png")
    umap_side = find_files(tmpdir, "ATAC_UMAP_sampleBySide.*.png")
    umap_qc = find_files(tmpdir, "ATAC_UMAP_withQC.*.png")
    cellcount_tissue = find_files(tmpdir, "ATAC_cellCount_perTissueByDonor.*.png")
    cellcount_donor = find_files(tmpdir, "ATAC_cellCount_perDonorPerTissue.*.png")

    if any([umap_sample, umap_donor, umap_side, umap_qc]):
        add_section_slide(prs, "Step 6: Integration & Clustering",
                          "Joint embedding and UMAP visualization")

        # UMAP by sample
        for png in umap_sample:
            add_image_slide(prs, png, "UMAP Colored by Sample")

        # UMAP by donor
        for png in umap_donor:
            add_image_slide(prs, png, "UMAP Colored by Donor")

        # Cell count barplots
        for png in cellcount_tissue:
            add_image_slide(prs, png, "Cell Count per Tissue by Donor")
        for png in cellcount_donor:
            add_image_slide(prs, png, "Cell Count per Donor per Tissue")

        # UMAP sample-by-side (multi-panel)
        for png in umap_side:
            add_image_slide(prs, png, "UMAP: Individual Samples")

        # Per-sample UMAP with QC metrics
        if umap_qc:
            overview_qc, persample_qc = group_by_sample(umap_qc, sample_ids)
            for png in overview_qc:
                add_image_slide(prs, png, short_name(png))
            if persample_qc:
                add_section_slide(prs, "Per-Sample QC UMAPs",
                                  "TSS enrichment and fragment count")
                for sid in sorted(persample_qc):
                    for png in persample_qc[sid]:
                        add_image_slide(prs, png, f"QC UMAP: {sid}")

    # Save
    out_pptx = os.path.join(tmpdir, f"{tissue_name}_ATAC_QC_Report.pptx")
    prs.save(out_pptx)
    print(f"[INFO] Saved presentation: {out_pptx}")
    return out_pptx


# ---------------------------------------------------------------------------
# Entry-point flows
# ---------------------------------------------------------------------------

def process_gcs(gcs_path, tissue=None):
    """GCS mode: download -> build -> upload."""
    gcs_base = gcs_path.rstrip("/") + "/"
    tissue_name = gcs_base.rstrip("/").split("/")[-1]
    if tissue:
        tissue_std = standardize_tissue_name(tissue)
    else:
        tissue_std = tissue_name.rsplit("_v", 1)[0]

    with tempfile.TemporaryDirectory() as tmpdir:
        print(f"[INFO] Downloading figures from {gcs_base}...")
        download_figures_and_stats(gcs_base, tmpdir)

        # Infer sample IDs from downloaded filenames
        sample_ids = extract_sample_ids_from_files(tmpdir, tissue_std)
        print(f"[INFO] Inferred {len(sample_ids)} sample IDs from filenames")

        pptx_path = build_presentation(
            tmpdir, tissue_name, tissue_std, sample_ids,
            source_label=gcs_base,
        )

        # Upload back to GCS
        gcs_dest = gcs_base + f"{tissue_name}_ATAC_QC_Report.pptx"
        print(f"[INFO] Uploading to {gcs_dest}...")
        subprocess.run(["gsutil", "-q", "cp", pptx_path, gcs_dest], check=True)
        print(f"[DONE] {gcs_dest}")


def process_gcs_parent(gcs_path):
    """Process all tissue subdirectories under a GCS parent path."""
    gcs_path = gcs_path.rstrip("/") + "/"
    contents = gsutil_ls(gcs_path)
    subdirs = [c for c in contents if c.endswith("/")]

    for tissue_dir in sorted(subdirs):
        try:
            process_gcs(tissue_dir)
        except Exception as e:
            print(f"[ERROR] Failed for {tissue_dir}: {e}")
            continue


def process_local(config_path, runtag, output_path=None):
    """Local mode: collect from workdir -> build -> save locally."""
    config = load_config(config_path)
    tissue = config['params']['tissue']
    suffix = config['params'].get('suffix') or standardize_tissue_name(tissue)
    workdir = config['paths']['workdir'].rstrip("/")

    tissue_std = suffix
    tissue_name = f"{tissue_std}_{runtag}"

    # Extract sample IDs from metadata
    try:
        sample_ids = extract_sample_ids_from_config(config, tissue_std)
        print(f"[INFO] Found {len(sample_ids)} sample IDs from metadata")
    except Exception as e:
        print(f"[WARN] Could not read sample metadata: {e}")
        sample_ids = set()

    with tempfile.TemporaryDirectory() as tmpdir:
        print(f"[INFO] Collecting figures from {workdir}...")
        collect_local_figures(workdir, runtag, tmpdir)

        # If we didn't get sample IDs from config, try inferring from filenames
        if not sample_ids:
            sample_ids = extract_sample_ids_from_files(tmpdir, tissue_std)
            if sample_ids:
                print(f"[INFO] Inferred {len(sample_ids)} sample IDs from filenames")

        pptx_path = build_presentation(
            tmpdir, tissue_name, tissue_std, sample_ids,
            source_label=workdir,
        )

        # Copy PPTX to final destination
        if output_path is None:
            output_path = os.path.join(workdir, f"{tissue_name}_ATAC_QC_Report.pptx")
        shutil.copy2(pptx_path, output_path)
        print(f"[DONE] {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Generate ATAC QC PowerPoint reports from pipeline outputs.",
        epilog="Examples:\n"
               "  %(prog)s config.yaml --runtag v0             # local pipeline run\n"
               "  %(prog)s --gcs gs://bucket/atacqc/Tissue_v0/   # GCS / Terra run\n",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("config", nargs="?",
                        help="Pipeline config YAML (for local runs)")
    parser.add_argument("--runtag", default="terra_run",
                        help="Run tag (default: terra_run)")
    parser.add_argument("--gcs",
                        help="GCS path to tissue directory or parent "
                             "(e.g. gs://bucket/atacqc/ for all)")
    parser.add_argument("--tissue",
                        help="Tissue name (e.g. 'Liver - Left Lobe'). "
                             "Overrides auto-detection from GCS path.")
    parser.add_argument("--output",
                        help="Output PPTX path (default: auto)")
    args = parser.parse_args()

    if args.gcs:
        gcs_path = args.gcs.rstrip("/") + "/"
        contents = gsutil_ls(gcs_path)
        subdirs = [c for c in contents if c.endswith("/")]
        has_steps = any(re.match(r'^\d+_', os.path.basename(d.rstrip("/")))
                        for d in subdirs)
        if has_steps:
            process_gcs(gcs_path, tissue=args.tissue)
        else:
            process_gcs_parent(gcs_path)
    elif args.config:
        process_local(args.config, args.runtag, args.output)
    else:
        parser.error("Provide either a config YAML (positional) or --gcs path")


if __name__ == "__main__":
    main()
