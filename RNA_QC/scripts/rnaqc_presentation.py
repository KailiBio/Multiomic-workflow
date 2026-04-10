#!/usr/bin/env python3
"""
Generate a PowerPoint QC presentation from RNA QC pipeline outputs.

Supports two input modes:

  Local (pipeline config YAML):
    python rnaqc_presentation.py config.yaml --runtag v0

  GCS (Terra / cloud):
    python rnaqc_presentation.py --gcs gs://bucket/rnaqc/Tissue_v0/
"""

import argparse
import os
import re
import shutil
import subprocess
import tempfile
import glob as globmod

import pandas as pd

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from PIL import Image
except ImportError as e:
    raise SystemExit(
        "Error: python-pptx and its dependencies (including Pillow) are required.\n"
        "    pip install python-pptx Pillow\n"
        f"Original error: {e}"
    )

# Import pipeline utilities when available (for local mode)
try:
    from rna_qc.utils import load_config, standardize_tissue_name
except ImportError:
    # Fallback implementations for standalone use / GCS-only mode
    try:
        import yaml
    except ImportError:
        raise SystemExit(
            "Error: PyYAML is required for standalone mode.\n"
            "    pip install pyyaml"
        )

    def load_config(config_path):
        if not os.path.exists(config_path):
            raise FileNotFoundError(f"Config file not found: {config_path}")
        with open(config_path, "r") as f:
            return yaml.safe_load(f)

    def standardize_tissue_name(tissue):
        tissue = re.sub(r'[()]', '', tissue)
        tissue = tissue.replace('-', '_')
        tissue = re.sub(r'\s+', '_', tissue)
        tissue = re.sub(r'_+', '_', tissue)
        return tissue.strip('_')

# ---------------------------------------------------------------------------
# Slide dimensions (widescreen 16:9)
# ---------------------------------------------------------------------------
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


# ---------------------------------------------------------------------------
# Sample-ID helpers
# ---------------------------------------------------------------------------

def is_per_sample(filepath, sample_ids):
    """Check if a file is per-sample (contains a sample ID in its name)."""
    base = os.path.basename(filepath)
    return any(sid in base for sid in sample_ids)


def extract_sample_ids_from_config(config, tissue_std):
    """Extract sample IDs (rnaID column) from the pipeline's sample_metadata TSV.

    Filters rows by matching standardized tissue name against ``tissue_std``.
    """
    meta_path = config['paths']['sample_metadata']
    tissue_config = config['params'].get('tissue', '---')
    df = pd.read_csv(meta_path, sep='\t', header=None, index_col=False,
                     names=["rnaID", "atacID", "species", "donorID",
                            "ageGroup", "gender", "tissue"])
    if tissue_config != "---":
        df = df[
            df["tissue"].apply(lambda x: standardize_tissue_name(str(x))) == tissue_std
        ]
    return set(df['rnaID'].dropna().tolist())


def extract_sample_ids_from_files(tmpdir, tissue_std):
    """Infer sample IDs from downloaded filenames (GCS fallback).

    Looks at scatter and doublet files which contain
    ``{tissue_std}-{sampleID}`` in their names.
    """
    sample_ids = set()
    # tissue_std uses underscores (never hyphens), so the first '-' after it
    # is always the sample-ID delimiter.  Require '.' or start-of-string
    # before the tissue name to avoid substring false-positives.
    pattern = re.compile(r'(?:^|\.)' + re.escape(tissue_std) + r'-([^.]+)')
    # Metadata field names that appear in filenames but are not sample IDs
    _metadata_fields = {'sampleID', 'donorID', 'aliquotID', 'species',
                        'tissue', 'ageGroup', 'gender'}
    for root, _dirs, files in os.walk(tmpdir):
        for f in files:
            m = pattern.search(f)
            if m and m.group(1) not in _metadata_fields:
                sample_ids.add(m.group(1))
    return sample_ids


# ---------------------------------------------------------------------------
# GCS helpers
# ---------------------------------------------------------------------------

def gsutil_ls(gcs_path):
    """List GCS path contents.

    Raises RuntimeError if gsutil is not available or the command fails.
    """
    try:
        result = subprocess.run(
            ["gsutil", "ls", gcs_path],
            capture_output=True, text=True
        )
    except FileNotFoundError as e:
        raise RuntimeError(
            "Failed to run 'gsutil ls'. Ensure gsutil is installed and on your PATH."
        ) from e

    if result.returncode != 0:
        stderr = result.stderr.strip()
        msg = f"'gsutil ls' failed for '{gcs_path}' (exit code {result.returncode})."
        if stderr:
            msg += f" stderr: {stderr}"
        raise RuntimeError(msg)

    return [line.strip() for line in result.stdout.strip().split("\n") if line.strip()]


def download_figures_and_stats(gcs_base, tmpdir):
    """Download all PNGs and TXT files from a tissue GCS directory."""
    steps = {
        "0_cellranger_qc":          ["*.png"],
        "2_pre_qc/figures":         ["*.png"],
        "3_filter_qc/figures":      ["*.png", "*.txt", "*.tsv"],
        "4_batch_correction/figures": ["*.png"],
        "5_clustering/figures":     ["*.png"],
        "5_clustering":             ["*.txt"],
    }
    for step, exts in steps.items():
        step_dir = os.path.join(tmpdir, step.replace("/", "_"))
        os.makedirs(step_dir, exist_ok=True)
        for ext in exts:
            result = subprocess.run(
                ["gsutil", "-m", "-q", "cp",
                 gcs_base + step + "/" + ext, step_dir + "/"],
                capture_output=True, text=True
            )
            if result.returncode != 0:
                print(f"[WARNING] gsutil cp failed for {gcs_base}{step}/{ext} (may not exist yet)")
    return tmpdir


# ---------------------------------------------------------------------------
# Local directory collection
# ---------------------------------------------------------------------------

def collect_local_figures(workdir, runtag, tmpdir):
    """Copy local pipeline outputs into a normalized tmpdir structure.

    The resulting directory layout matches what ``download_figures_and_stats``
    produces so that ``find_files`` + ``build_presentation`` work identically
    for both GCS and local runs.
    """
    dir_mapping = {
        "cellranger_qc": "0_cellranger_qc",
        "pre_qc_assessment": "2_pre_qc_figures",
        f"filter_qc.{runtag}": "3_filter_qc_figures",
        f"batch_correction.{runtag}": "4_batch_correction_figures",
        f"clustering_and_qc_reassessment.{runtag}": "5_clustering_figures",
    }
    found_any = False
    for local_name, tmp_name in dir_mapping.items():
        src = os.path.join(workdir, local_name)
        if os.path.isdir(src):
            dst = os.path.join(tmpdir, tmp_name)
            shutil.copytree(src, dst, dirs_exist_ok=True)
            found_any = True
            print(f"[INFO]   {local_name}/ -> {tmp_name}/")
    if not found_any:
        print(f"[WARN] No pipeline output directories found in {workdir}")
    return tmpdir


# ---------------------------------------------------------------------------
# File finding / naming helpers
# ---------------------------------------------------------------------------

def find_files(base_dir, pattern):
    """Find files matching pattern recursively, deduplicating by filename.

    Dedup-by-filename is intentional: the normalized tmpdir has distinct
    subdirectories for each step, and duplicate basenames would indicate the
    same file copied to multiple stage dirs (which we want to show only once).
    """
    matches = []
    seen = set()
    for root, dirs, files in os.walk(base_dir):
        for f in files:
            if globmod.fnmatch.fnmatch(f, pattern) and f not in seen:
                matches.append(os.path.join(root, f))
                seen.add(f)
    matches.sort()
    return matches


def short_name(filepath):
    """Extract a readable name from the file path."""
    name = os.path.basename(filepath)
    name = name.replace(".png", "")
    return name


def group_by_sample(files, sample_ids):
    """Separate files into overview and per-sample lists."""
    overview = []
    per_sample = {}
    for f in files:
        base = os.path.basename(f)
        matched_id = None
        for sid in sample_ids:
            if sid in base:
                matched_id = sid
                break
        if matched_id:
            per_sample.setdefault(matched_id, []).append(f)
        else:
            overview.append(f)
    return overview, per_sample


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def add_title_slide(prs, tissue_name, source_label):
    """Add title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"RNA QC Report: {tissue_name.replace('_', ' ')}"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    p.alignment = PP_ALIGN.CENTER
    if source_label:
        p2 = tf.add_paragraph()
        p2.text = f"Source: {source_label}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p2.alignment = PP_ALIGN.CENTER


def add_section_slide(prs, title, subtitle=""):
    """Add a section divider slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    p.alignment = PP_ALIGN.CENTER
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(16)
        p2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p2.alignment = PP_ALIGN.CENTER


def add_image_slide(prs, image_path, title="", max_height=Inches(6.2)):
    """Add a slide with a single image, scaled to fit."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if title:
        txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(12.5), Inches(0.6))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

    with Image.open(image_path) as img:
        img_w, img_h = img.size

    max_width = Inches(12.5)
    top_margin = Inches(0.8) if title else Inches(0.3)

    scale_w = max_width / Emu(int(img_w * 914400 / 96))
    scale_h = max_height / Emu(int(img_h * 914400 / 96))
    scale = min(scale_w, scale_h, 1.0)

    final_w = int(img_w * 914400 / 96 * scale)
    final_h = int(img_h * 914400 / 96 * scale)

    left = int((SLIDE_WIDTH - final_w) / 2)
    slide.shapes.add_picture(image_path, left, int(top_margin), final_w, final_h)


def add_two_images_slide(prs, img1, img2, title=""):
    """Add a slide with two images side by side."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    if title:
        txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), Inches(12.5), Inches(0.6))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER

    top = Inches(0.8) if title else Inches(0.3)
    max_w = Inches(6.2)
    max_h = Inches(6.0)

    for i, img_path in enumerate([img1, img2]):
        with Image.open(img_path) as img:
            img_w, img_h = img.size
        scale_w = max_w / Emu(int(img_w * 914400 / 96))
        scale_h = max_h / Emu(int(img_h * 914400 / 96))
        scale = min(scale_w, scale_h, 1.0)
        final_w = int(img_w * 914400 / 96 * scale)
        final_h = int(img_h * 914400 / 96 * scale)
        left = Inches(0.3) if i == 0 else Inches(6.8)
        slide.shapes.add_picture(img_path, int(left), int(top), final_w, final_h)


def add_text_slide(prs, title, text_content):
    """Add a slide with text content (for stats files)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(12), Inches(6))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    lines = text_content.strip().split("\n")
    for i, line in enumerate(lines):
        if i == 0:
            p2 = tf2.paragraphs[0]
        else:
            p2 = tf2.add_paragraph()
        p2.text = line
        p2.font.size = Pt(12)
        p2.font.name = "Courier New"


def add_two_column_text_slide(prs, title, text_content):
    """Add a slide with text split into two columns by section separators."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(12.5), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    sections = []
    current = []
    for line in text_content.strip().split("\n"):
        if line.strip().startswith("---"):
            if current:
                sections.append("\n".join(current))
                current = []
        else:
            current.append(line)
    if current:
        sections.append("\n".join(current))

    mid = (len(sections) + 1) // 2
    left_text = "\n\n".join(sections[:mid])
    right_text = "\n\n".join(sections[mid:])

    for col_idx, col_text in enumerate([left_text, right_text]):
        if not col_text.strip():
            continue
        left = Inches(0.3) if col_idx == 0 else Inches(6.8)
        tb = slide.shapes.add_textbox(left, Inches(1.0), Inches(6.0), Inches(6.0))
        tf2 = tb.text_frame
        tf2.word_wrap = True
        lines = col_text.split("\n")
        for i, line in enumerate(lines):
            if i == 0:
                p2 = tf2.paragraphs[0]
            else:
                p2 = tf2.add_paragraph()
            p2.text = line
            p2.font.size = Pt(11)
            p2.font.name = "Courier New"


# ---------------------------------------------------------------------------
# Main presentation builder
# ---------------------------------------------------------------------------

def build_presentation(tmpdir, tissue_name, tissue_std, sample_ids, source_label=""):
    """Build the PowerPoint presentation from figures in tmpdir.

    Parameters
    ----------
    tmpdir : str
        Temporary directory containing normalized figure subdirectories.
    tissue_name : str
        Display name (e.g. "Kidney_Left_v0").
    tissue_std : str
        Standardized tissue name for filename matching (e.g. "Kidney_Left").
    sample_ids : set[str]
        Set of sample ID strings used to distinguish per-sample vs overview.
    source_label : str
        Shown on title slide (GCS path or local workdir path).
    """
    print(f"[INFO] Building presentation for {tissue_name}...")
    if sample_ids:
        print(f"[INFO]   {len(sample_ids)} sample IDs for per-sample detection")
    else:
        print("[WARN]   No sample IDs found — all plots will be treated as overview")

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    add_title_slide(prs, tissue_std, source_label)

    # --- Step 0: CellRanger QC ---
    cr_pngs = find_files(tmpdir, "CellRanger_*.png")
    if cr_pngs:
        add_section_slide(prs, "Step 0: CellRanger QC Summary")
        for png in cr_pngs:
            add_image_slide(prs, png, short_name(png))

    # --- Step 2: Pre-QC Assessment ---
    preqc_pngs = find_files(tmpdir, "RNA_QC_violin.*.png")
    preqc_cumul = find_files(tmpdir, "RNA_QC_cumulative.*.png")
    preqc_scatter = find_files(tmpdir, "*cellCounts_geneCounts_scatter.png")
    preqc_doublet = [f for f in find_files(tmpdir, "RNA_QC_doublet_hist.*.png")
                     if "filterqc" not in os.path.basename(f) and "3_filter" not in f]
    preqc_highest = [f for f in find_files(tmpdir, "highest_expr_genes.*.png")
                     if "postFilter" not in os.path.basename(f)]
    preqc_umap = find_files(tmpdir, "umap.*BeforeFiltering*.png")

    if any([preqc_pngs, preqc_cumul, preqc_scatter]):
        add_section_slide(prs, "Step 2: Pre-QC Assessment", "QC metrics before filtering")

        # Highest expressed genes
        for png in preqc_highest:
            add_image_slide(prs, png, "Highest Expressed Genes (Pre-Filter)")

        # Overview scatter
        overview_scatter = [f for f in preqc_scatter
                           if not is_per_sample(f, sample_ids)]
        for png in overview_scatter:
            add_image_slide(prs, png, "Cell Counts vs Gene Counts")

        # Violin plots
        violin_metrics = {}
        for f in preqc_pngs:
            base = os.path.basename(f)
            metric = base.split(".")[-2] if "." in base else base
            violin_metrics.setdefault(metric, []).append(f)
        for metric in sorted(violin_metrics):
            for png in violin_metrics[metric]:
                add_image_slide(prs, png, f"Violin: {metric}")

        # Cumulative plots
        for png in preqc_cumul:
            metric = os.path.basename(png).split(".")[-2]
            add_image_slide(prs, png, f"Cumulative Distribution: {metric}")

        # Doublet histograms
        if preqc_doublet:
            add_section_slide(prs, "Pre-QC: Doublet Detection")
            for png in preqc_doublet:
                add_image_slide(prs, png, short_name(png))

        # Pre-QC UMAPs
        overview_umap = [f for f in preqc_umap
                         if not is_per_sample(f, sample_ids)]
        sample_umap = [f for f in preqc_umap
                       if is_per_sample(f, sample_ids)]
        for png in overview_umap:
            add_image_slide(prs, png, "UMAP: Leiden Clusters (Before Filtering)")
        if sample_umap:
            for i in range(0, len(sample_umap), 2):
                if i + 1 < len(sample_umap):
                    add_two_images_slide(prs, sample_umap[i], sample_umap[i + 1],
                                         "Pre-QC UMAP by Sample")
                else:
                    add_image_slide(prs, sample_umap[i], short_name(sample_umap[i]))

        # Per-sample scatters
        sample_scatter = [f for f in preqc_scatter
                          if is_per_sample(f, sample_ids)]
        if sample_scatter:
            for i in range(0, len(sample_scatter), 2):
                if i + 1 < len(sample_scatter):
                    add_two_images_slide(prs, sample_scatter[i], sample_scatter[i + 1],
                                         "Per-Sample: Cell vs Gene Counts")
                else:
                    add_image_slide(prs, sample_scatter[i], short_name(sample_scatter[i]))

    # --- Step 3: Filter QC ---
    filter_stats = find_files(tmpdir, "*filtering_stat_counts.txt")
    upset_all = find_files(tmpdir, "QC_filtering_upset.all.*.png")
    upset_filt = find_files(tmpdir, "QC_filtering_upset.filteringOnly.*.png")
    filter_doublet = find_files(tmpdir, "RNA_QC_doublet_hist.filterqc.*.png")
    hvg = find_files(tmpdir, "*highlyVariableGenes.filterqc.*.png")
    pca_scatter = find_files(tmpdir, "*pca_scatter.filterqc.*.png")
    pca_var = find_files(tmpdir, "*var_ratio.filterqc.*.png")
    filter_umap = find_files(tmpdir, "umap.filterqc.*.png")

    if any([filter_stats, upset_all, filter_umap]):
        add_section_slide(prs, "Step 3: QC Filtering", "Cell and gene filtering results")

        # Filtering stats
        for stat_file in filter_stats:
            with open(stat_file) as f:
                add_text_slide(prs, "Filtering Statistics", f.read())

        # Upset plots - pair all vs filtering-only per sample
        samples_upset = sorted(set(
            os.path.basename(f).split(".")[-2] for f in upset_all
        ))
        for sample in samples_upset:
            all_f = [f for f in upset_all if sample in os.path.basename(f)]
            filt_f = [f for f in upset_filt if sample in os.path.basename(f)]
            if all_f and filt_f:
                add_two_images_slide(prs, all_f[0], filt_f[0],
                                     f"Upset: {sample}  (Left: all metrics incl. defaults | Right: explicit cutoffs only)")
            elif all_f:
                add_image_slide(prs, all_f[0], f"Upset (All QC metrics incl. defaults): {sample}")

        # Doublet histograms (with applied cutoff)
        if filter_doublet:
            for png in filter_doublet:
                add_image_slide(prs, png, short_name(png))

        # HVG, PCA
        for png in hvg:
            add_image_slide(prs, png, "Highly Variable Genes (Post-Filter)")
        if pca_scatter and pca_var:
            add_two_images_slide(prs, pca_scatter[0], pca_var[0],
                                 "PCA: Scatter & Variance Ratio")
        else:
            for png in pca_scatter + pca_var:
                add_image_slide(prs, png, short_name(png))

        # Filter UMAPs
        for png in filter_umap:
            add_image_slide(prs, png, short_name(png))

    # --- Step 4: Batch Correction ---
    harmony_pngs = find_files(tmpdir, "*Harmony*.png")
    harmony_seen = set(os.path.basename(f) for f in harmony_pngs)
    harmony_pngs += [f for f in find_files(tmpdir, "*batchCorrection*.png")
                     if os.path.basename(f) not in harmony_seen]
    if harmony_pngs:
        add_section_slide(prs, "Step 4: Batch Correction", "Harmony integration")
        for png in harmony_pngs:
            add_image_slide(prs, png, "Batch Correction UMAP")

    # --- Step 5: Clustering ---
    cluster_stats = find_files(tmpdir, "*stat_counts*.txt")
    cluster_stats = [f for f in cluster_stats if "filtering_stat" not in os.path.basename(f)]
    cluster_umap = find_files(tmpdir, "umap.LeidenCluster*.png")
    cluster_check = find_files(tmpdir, "umap.check_leiden*.png")
    cluster_cellcount = find_files(tmpdir, "Number_cellbarcode*.png")
    # Match any key: sampleID, donorID, aliquotID, etc.
    cluster_umap_sample = find_files(tmpdir, "UMAP_LeidenCluster_by*.*.png")
    qc_by_cluster = find_files(tmpdir, "QCmetric_bycluster.*.png")
    highest_post = find_files(tmpdir, "highest_expr_genes.postFilter.*.png")

    if any([cluster_umap, qc_by_cluster]):
        add_section_slide(prs, "Step 5: Clustering & QC Reassessment",
                          "Final clustering results")

        # Stats (2-column layout)
        for stat_file in cluster_stats:
            with open(stat_file) as f:
                add_two_column_text_slide(prs, "Cell Count Statistics", f.read())

        # Highest expressed genes post-filter
        for png in highest_post:
            add_image_slide(prs, png, "Highest Expressed Genes (Post-Filter)")

        # Leiden resolution check
        for png in cluster_check:
            add_image_slide(prs, png, "Leiden Resolution Comparison")

        # Overview UMAPs
        overview_cluster = [f for f in cluster_umap
                            if not is_per_sample(f, sample_ids)]
        for png in overview_cluster:
            add_image_slide(prs, png, short_name(png))

        # UMAP by sample/donor/aliquot
        for png in cluster_umap_sample:
            add_image_slide(prs, png, "UMAP: Leiden Clusters by Sample")

        # Cell count barplot
        for png in cluster_cellcount:
            add_image_slide(prs, png, "Cell Count per Cluster by Sample")

        # QC metrics by cluster
        qc_overview = [f for f in qc_by_cluster
                       if "sampleID" not in os.path.basename(f)]
        qc_persample = [f for f in qc_by_cluster
                        if "sampleID" in os.path.basename(f)]

        if qc_overview:
            add_section_slide(prs, "QC Metrics by Cluster")
            for png in qc_overview:
                metric = os.path.basename(png).split(".")[-2]
                add_image_slide(prs, png, f"QC by Cluster: {metric}")

        if qc_persample:
            add_section_slide(prs, "QC Metrics by Cluster (per Sample)")
            for png in qc_persample:
                metric = os.path.basename(png).split(".")[-2]
                add_image_slide(prs, png, f"QC by Cluster (per Sample): {metric}")

        # Per-sample UMAPs
        sample_cluster_umap = [f for f in cluster_umap
                               if is_per_sample(f, sample_ids)]
        if sample_cluster_umap:
            add_section_slide(prs, "Per-Sample UMAPs (Post-Clustering)")
            for i in range(0, len(sample_cluster_umap), 2):
                if i + 1 < len(sample_cluster_umap):
                    add_two_images_slide(prs, sample_cluster_umap[i],
                                         sample_cluster_umap[i + 1],
                                         "Per-Sample UMAP with QC Metrics")
                else:
                    add_image_slide(prs, sample_cluster_umap[i],
                                    short_name(sample_cluster_umap[i]))

    # Save
    out_pptx = os.path.join(tmpdir, f"{tissue_name}_RNA_QC_Report.pptx")
    prs.save(out_pptx)
    print(f"[INFO] Saved presentation: {out_pptx}")
    return out_pptx


# ---------------------------------------------------------------------------
# Entry-point flows
# ---------------------------------------------------------------------------

def process_gcs(gcs_path, tissue=None):
    """GCS mode: download -> build -> upload (backwards compatible)."""
    gcs_base = gcs_path.rstrip("/") + "/"
    tissue_name = gcs_base.rstrip("/").split("/")[-1]
    if tissue:
        tissue_std = standardize_tissue_name(tissue)
    else:
        tissue_std = tissue_name.rsplit("_v", 1)[0]

    with tempfile.TemporaryDirectory() as tmpdir:
        print(f"[INFO] Downloading figures from {gcs_base}...")
        download_figures_and_stats(gcs_base, tmpdir)

        # Infer sample IDs from downloaded filenames
        sample_ids = extract_sample_ids_from_files(tmpdir, tissue_std)

        # If no sample IDs found, try progressively shorter tissue prefixes.
        # Handles folder names like "Brain_Cerebellum_yale_replication" where
        # the filenames use just "Brain_Cerebellum" as the tissue prefix.
        if not sample_ids and '_' in tissue_std:
            candidate = tissue_std
            while '_' in candidate:
                candidate = candidate.rsplit('_', 1)[0]
                sample_ids = extract_sample_ids_from_files(tmpdir, candidate)
                if sample_ids:
                    print(f"[INFO] Matched tissue prefix '{candidate}' "
                          f"(folder: {tissue_name})")
                    tissue_std = candidate
                    break

        print(f"[INFO] Inferred {len(sample_ids)} sample IDs from filenames")
        for sid in sorted(sample_ids):
            print(f"[INFO]   - {sid}")

        pptx_path = build_presentation(
            tmpdir, tissue_name, tissue_std, sample_ids,
            source_label=gcs_base,
        )

        # Upload back to GCS
        gcs_dest = gcs_base + f"{tissue_name}_RNA_QC_Report.pptx"
        print(f"[INFO] Uploading to {gcs_dest}...")
        subprocess.run(["gsutil", "-q", "cp", pptx_path, gcs_dest], check=True)
        print(f"[DONE] {gcs_dest}")


def process_gcs_parent(gcs_path):
    """Process all tissue subdirectories under a GCS parent path."""
    gcs_path = gcs_path.rstrip("/") + "/"
    contents = gsutil_ls(gcs_path)
    subdirs = [c for c in contents if c.endswith("/")]

    for tissue_dir in sorted(subdirs):
        try:
            process_gcs(tissue_dir)
        except Exception as e:
            print(f"[ERROR] Failed for {tissue_dir}: {e}")
            continue


def process_local(config_path, runtag, output_path=None):
    """Local mode: collect from workdir -> build -> save locally."""
    config = load_config(config_path)
    tissue = config['params']['tissue']
    suffix = config['params'].get('suffix') or standardize_tissue_name(tissue)
    workdir = config['paths']['workdir'].rstrip("/")

    tissue_std = suffix  # standardized name for filename matching
    tissue_name = f"{tissue_std}_{runtag}"

    # Extract sample IDs from metadata
    try:
        sample_ids = extract_sample_ids_from_config(config, tissue_std)
        print(f"[INFO] Found {len(sample_ids)} sample IDs from metadata")
    except Exception as e:
        print(f"[WARN] Could not read sample metadata: {e}")
        sample_ids = set()

    with tempfile.TemporaryDirectory() as tmpdir:
        print(f"[INFO] Collecting figures from {workdir}...")
        collect_local_figures(workdir, runtag, tmpdir)

        # If we didn't get sample IDs from config, try inferring from filenames
        if not sample_ids:
            sample_ids = extract_sample_ids_from_files(tmpdir, tissue_std)
            if sample_ids:
                print(f"[INFO] Inferred {len(sample_ids)} sample IDs from filenames")

        pptx_path = build_presentation(
            tmpdir, tissue_name, tissue_std, sample_ids,
            source_label=workdir,
        )

        # Copy PPTX to final destination
        if output_path is None:
            output_path = os.path.join(workdir, f"{tissue_name}_RNA_QC_Report.pptx")
        shutil.copy2(pptx_path, output_path)
        print(f"[DONE] {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Generate RNA QC PowerPoint reports from pipeline outputs.",
        epilog="Examples:\n"
               "  %(prog)s config.yaml --runtag v0          # local pipeline run\n"
               "  %(prog)s --gcs gs://bucket/rnaqc/Tissue_v0/  # GCS / Terra run\n",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("config", nargs="?",
                        help="Pipeline config YAML (for local runs)")
    parser.add_argument("--runtag", default="terra_run",
                        help="Run tag (default: terra_run)")
    parser.add_argument("--gcs",
                        help="GCS path to tissue directory or parent "
                             "(e.g. gs://bucket/rnaqc/ for all)")
    parser.add_argument("--tissue",
                        help="Tissue name (e.g. 'Liver - Left Lobe'). "
                             "Overrides auto-detection from GCS path.")
    parser.add_argument("--output",
                        help="Output PPTX path (default: auto)")
    args = parser.parse_args()

    if args.gcs:
        # GCS mode
        gcs_path = args.gcs.rstrip("/") + "/"
        contents = gsutil_ls(gcs_path)
        subdirs = [c for c in contents if c.endswith("/")]
        has_steps = any(re.match(r'^\d+_', os.path.basename(d.rstrip("/")))
                        for d in subdirs)
        if has_steps:
            process_gcs(gcs_path, tissue=args.tissue)
        else:
            process_gcs_parent(gcs_path)
    elif args.config:
        # Local mode
        process_local(args.config, args.runtag, args.output)
    else:
        parser.error("Provide either a config YAML (positional) or --gcs path")


if __name__ == "__main__":
    main()
